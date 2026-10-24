# Copyright 2026 TrillionniumFoundation; SPDX-License-Identifier: Apache-2.0
"""Actual OOXML and independent reopen checks; no ExcelJS/DC execution."""
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'deploy'))
from dc_compat.filesystem import FileManager

DOCUMENT_PYTHON = '/home/qian-qi/.local/lib/hepta-runtime/document-env-formats-20261006/bin/python'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


class Config:
    def __init__(self, root): self.workspace, self.document_python = root, DOCUMENT_PYTHON
    def get(self, key, default=None): return [str(self.workspace)] if key == 'allowedDirectories' else default
    def resolve_path(self, raw):
        path = Path(raw)
        path = (path if path.is_absolute() else self.workspace / path).resolve()
        if path != self.workspace and self.workspace not in path.parents: raise ValueError('outside fixture')
        return path


class OfficeObjects(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='hepta-office-objects-')
        self.root = Path(self.temporary.name)
        self.manager = FileManager(Config(self.root))

    async def asyncTearDown(self): self.temporary.cleanup()

    async def call(self, name, arguments):
        result = await self.manager.invoke(name, arguments)
        self.assertFalse(result.get('isError'), result)
        return result

    async def write(self, content, name='fixture.xlsx', mode=None):
        arguments = {'path': name, 'content': json.dumps(content, ensure_ascii=False)}
        if mode: arguments['mode'] = mode
        await self.call('write_file', arguments)
        return self.root / name

    async def edit(self, region, content, name='fixture.xlsx'):
        return await self.call('edit_block', {'file_path': name, 'range': region, 'content': content})

    def xml(self, path, member='xl/worksheets/sheet1.xml'):
        from lxml import etree
        with zipfile.ZipFile(path) as archive: return etree.fromstring(archive.read(member))

    def workbook(self, path, **options):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(path.read_bytes()), rich_text=True, **options)

    async def test_new_formula_objects_keep_numeric_zero_boolean_string_error_and_missing_caches(self):
        values = [42, 0, False, True, '缓存中文', {'error': '#DIV/0!'}, None]
        path = await self.write([[{'formula': '1+1', 'result': value} for value in values]])
        root = self.xml(path)
        cells = root.findall('.//' + S + 'c')
        self.assertEqual(len(cells), len(values))
        for cell in cells: self.assertEqual(cell.find(S + 'f').text, '1+1')
        self.assertEqual([cell.get('t') for cell in cells], [None, None, 'b', 'b', 'str', 'e', None])
        self.assertEqual([cell.findtext(S + 'v') for cell in cells], ['42', '0', '0', '1', '缓存中文', '#DIV/0!', None])
        book = self.workbook(path, data_only=True)
        self.assertEqual([cell.value for cell in book.active[1]], [42, 0, False, True, '缓存中文', '#DIV/0!', None]); book.close()
        book = self.workbook(path, data_only=False)
        self.assertEqual(book.active['A1'].value, '=1+1'); book.close()
        result = await self.call('read_file', {'path': path.name})
        projected = json.loads(result['content'][0]['text'].split('\n\n', 1)[1])
        self.assertEqual(projected[0][:5], [42, 0, False, True, '缓存中文'])
        self.assertEqual(projected[0][5], {'error': '#DIV/0!'})

    async def test_richtext_keeps_every_exceljs_font_field_color_and_run_boundaries(self):
        font = {'name': 'Arial', 'charset': 1, 'family': 2, 'bold': True, 'italic': True, 'strike': True,
                'outline': True, 'shadow': True, 'condense': True, 'extend': True, 'size': 14,
                'underline': 'doubleAccounting', 'vertAlign': 'superscript', 'scheme': 'minor',
                'color': {'argb': 'FF123456'}}
        rich = {'richText': [{'font': font, 'text': '  格式一 '}, {'text': 'plain'},
                            {'font': {'color': {'theme': 4, 'tint': .25}, 'underline': True}, 'text': ' theme'},
                            {'font': {'color': {'indexed': 7}}, 'text': ' indexed'}]}
        path = await self.write([[rich]])
        root = self.xml(path)
        runs = root.findall('.//' + S + 'r')
        self.assertEqual([run.findtext(S + 't') for run in runs], ['  格式一 ', 'plain', ' theme', ' indexed'])
        properties = runs[0].find(S + 'rPr')
        self.assertEqual(properties.find(S + 'rFont').get('val'), 'Arial')
        self.assertEqual(properties.find(S + 'sz').get('val'), '14')
        self.assertEqual(properties.find(S + 'u').get('val'), 'doubleAccounting')
        self.assertEqual(properties.find(S + 'color').get('rgb'), 'FF123456')
        for tag in ('b', 'i', 'strike', 'outline', 'shadow', 'condense', 'extend'):
            self.assertIsNotNone(properties.find(S + tag))
        self.assertEqual(runs[2].find(S + 'rPr/' + S + 'color').get('theme'), '4')
        self.assertEqual(runs[2].find(S + 'rPr/' + S + 'color').get('tint'), '0.25')
        self.assertEqual(runs[3].find(S + 'rPr/' + S + 'color').get('indexed'), '7')
        book = self.workbook(path)
        cell = book.active['A1'].value
        self.assertEqual(str(cell), '  格式一 plain theme indexed')
        self.assertEqual(cell[0].font.rFont, 'Arial'); self.assertEqual(cell[0].font.u, 'doubleAccounting')
        self.assertTrue(cell[0].font.b); self.assertTrue(cell[0].font.i)
        book.close()

    async def test_external_and_internal_hyperlinks_preserve_targets_tooltips_and_display_text(self):
        external = {'text': '链接显示', 'hyperlink': 'https://fixture.invalid/path?q=甲&next=乙', 'tooltip': '说明中文'}
        internal = {'text': 'Jump', 'hyperlink': 'Other!A1', 'tooltip': 'Internal'}
        path = await self.write({'Sheet1': [[external, internal]], 'Other': [['target']]})
        root = self.xml(path)
        links = root.find(S + 'hyperlinks')
        self.assertEqual(len(links), 2)
        self.assertEqual(links[1].get('location'), 'Other!A1')
        rels = self.xml(path, 'xl/worksheets/_rels/sheet1.xml.rels')
        targets = {item.get('Id'): item.get('Target') for item in rels}
        self.assertEqual(targets[links[0].get(R + 'id')], external['hyperlink'])
        self.assertEqual(links[0].get('tooltip'), external['tooltip'])
        book = self.workbook(path)
        self.assertEqual(book['Sheet1']['A1'].hyperlink.target, external['hyperlink'])
        self.assertEqual(book['Sheet1']['A1'].hyperlink.tooltip, external['tooltip'])
        self.assertEqual(book['Sheet1']['A1'].value, external['text']); book.close()

    async def test_edit_objects_preserves_other_cells_cache_styles_and_every_unrelated_zip_member(self):
        import openpyxl
        from openpyxl.styles import PatternFill
        book = openpyxl.Workbook(); book.active['B2'] = 'before'
        book.active['B2'].fill = PatternFill('solid', fgColor='FFFF0000')
        book.create_sheet('Other')['C3'] = 'UNCHANGED'
        path = self.root / 'fixture.xlsx'; book.save(path); book.close()
        await self.edit('Sheet!A1', [[{'formula': '1+1', 'result': 2}]])
        with zipfile.ZipFile(path, 'a') as archive: archive.writestr('opaque/custom.bin', b'original-opaque-bytes')
        with zipfile.ZipFile(path) as archive: before = {name: archive.read(name) for name in archive.namelist()}
        await self.edit('Sheet!B2', [[{'richText': [{'font': {'bold': True}, 'text': 'BOLD'}]},
                                   {'text': 'Link', 'hyperlink': 'https://fixture.invalid/owned'},
                                   {'formula': '"cached"', 'result': '缓存'}, {'formula': '1=2', 'result': False}]])
        with zipfile.ZipFile(path) as archive:
            for name, raw in before.items():
                if name != 'xl/worksheets/sheet1.xml': self.assertEqual(archive.read(name), raw, name)
        book = self.workbook(path)
        self.assertEqual(book['Sheet']['B2'].fill.fgColor.rgb, 'FFFF0000')
        self.assertTrue(book['Sheet']['B2'].value[0].font.b)
        self.assertEqual(book['Other']['C3'].value, 'UNCHANGED'); book.close()
        cached = self.workbook(path, data_only=True)
        self.assertEqual(cached['Sheet']['A1'].value, 2)
        self.assertEqual(cached['Sheet']['D2'].value, '缓存')
        self.assertIs(cached['Sheet']['E2'].value, False); cached.close()

    async def test_append_objects_and_new_sheet_are_atomic_across_all_sheets(self):
        path = await self.write({'Sheet1': [['first']], 'Untouched': [['original']]})
        await self.write({'Sheet1': [[{'formula': '2+3', 'result': 5}]],
                          'Added': [[{'richText': [{'font': {'italic': True}, 'text': 'run'}]},
                                     {'text': 'new', 'hyperlink': 'https://fixture.invalid/new'}]]}, mode='append')
        cached = self.workbook(path, data_only=True)
        self.assertEqual(cached['Sheet1']['A2'].value, 5)
        self.assertEqual(cached['Untouched']['A1'].value, 'original'); cached.close()
        book = self.workbook(path)
        self.assertTrue(book['Added']['A1'].value[0].font.i)
        self.assertEqual(book['Added']['B1'].hyperlink.target, 'https://fixture.invalid/new'); book.close()
        before = path.read_bytes()
        bad = await self.manager.invoke('write_file', {'path': path.name, 'mode': 'append',
            'content': json.dumps({'Sheet1': [['MUST-NOT-COMMIT']], 'Added': [[{'richText': [{'font': {'unknown': 1}, 'text': 'bad'}]}]]})})
        self.assertTrue(bad.get('isError')); self.assertEqual(path.read_bytes(), before)

    async def test_replacing_link_removes_only_its_relationship_and_keeps_other_links(self):
        path = await self.write([[{'text': 'first', 'hyperlink': 'https://fixture.invalid/first'},
                                 {'text': 'second', 'hyperlink': 'https://fixture.invalid/second'}]])
        await self.edit('Sheet1!A1', [['ordinary']])
        root = self.xml(path); links = root.find(S + 'hyperlinks')
        self.assertEqual([link.get('ref') for link in links], ['B1'])
        rels = self.xml(path, 'xl/worksheets/_rels/sheet1.xml.rels')
        self.assertEqual([item.get('Target') for item in rels], ['https://fixture.invalid/second'])
        await self.edit('Sheet1!A1', [[{'text': 'replacement', 'hyperlink': 'https://fixture.invalid/replaced'}]])
        book = self.workbook(path)
        self.assertEqual(book.active['A1'].hyperlink.target, 'https://fixture.invalid/replaced')
        self.assertEqual(book.active['B1'].hyperlink.target, 'https://fixture.invalid/second'); book.close()

    async def test_shared_and_array_formula_objects_reopen_with_translated_formulas_and_cache(self):
        rows = [[{'formula': 'ROW()', 'shareType': 'shared', 'ref': 'A1:A3', 'result': 1}, None,
                 {'formula': 'A1+1', 'result': 2}, None,
                 {'formula': 'SUM(A1:A3)', 'shareType': 'array', 'ref': 'E1:E3', 'result': 6}],
                [{'sharedFormula': 'A1', 'result': 2}, None, {'sharedFormula': 'C1', 'result': 3}],
                [{'sharedFormula': 'A1', 'result': 3}]]
        path = await self.write(rows)
        root = self.xml(path)
        self.assertEqual(root.find('.//' + S + 'c[@r="A1"]/' + S + 'f').get('ref'), 'A1:A3')
        book = self.workbook(path)
        self.assertEqual(book.active['A2'].value, '=ROW()')
        self.assertEqual(book.active['C2'].value, '=A2+1')
        self.assertEqual(book.active['E1'].value.ref, 'E1:E3'); book.close()
        cached = self.workbook(path, data_only=True)
        self.assertEqual([cached.active[f'A{index}'].value for index in (1, 2, 3)], [1, 2, 3])
        self.assertEqual(cached.active['E1'].value, 6); cached.close()
        before = path.read_bytes()
        refused = await self.manager.invoke('edit_block', {'file_path': path.name, 'range': 'Sheet1!A1', 'content': [['orphan-master']]})
        self.assertTrue(refused.get('isError')); self.assertEqual(path.read_bytes(), before)

    async def test_unsupported_object_data_refuses_before_publishing_instead_of_flattening(self):
        path = await self.write([['ORIGINAL']])
        original = path.read_bytes()
        bad_values = [{'richText': [{'font': {'unsupportedProperty': True}, 'text': 'LOST'}]},
                      {'sharedFormula': 'Z99', 'result': 2}, {'sharedString': 23},
                      {'formula': '1+1', 'result': {'richText': [{'text': 'bad'}]}}]
        for value in bad_values:
            reply = await self.manager.invoke('edit_block', {'file_path': path.name, 'range': 'Sheet1!A1', 'content': [[value]]})
            self.assertTrue(reply.get('isError'), reply); self.assertEqual(path.read_bytes(), original)

    async def test_exceljs_json_fallback_retains_arbitrary_objects_and_lists_as_json_strings(self):
        values = [{'arbitrary': {'汉字': True}, 'n': 0}, [1, 'two', None], {'text': 'only text without link'}]
        path = await self.write([values])
        book = self.workbook(path)
        self.assertEqual([json.loads(cell.value) for cell in book.active[1]], values); book.close()

    async def test_empty_formula_cache_and_font_color_precedence_are_not_silently_changed(self):
        path = await self.write([[{'formula': 'IF(TRUE,"",0)', 'result': ''},
                                 {'richText': [{'font': {'color': {}}, 'text': 'auto'},
                                               {'font': {'color': {'argb': 'FFAABBCC', 'theme': 4, 'indexed': 2}}, 'text': 'argb wins'}]}]])
        root = self.xml(path)
        colors = root.findall('.//' + S + 'rPr/' + S + 'color')
        self.assertEqual(colors[0].get('auto'), '1')
        self.assertEqual(colors[1].get('rgb'), 'FFAABBCC')
        self.assertIsNone(colors[1].get('theme'))
        result = await self.call('read_file', {'path': path.name})
        projected = json.loads(result['content'][0]['text'].split('\n\n', 1)[1])
        self.assertEqual(projected[0], ['', 'autoargb wins'])


if __name__ == '__main__': unittest.main()
