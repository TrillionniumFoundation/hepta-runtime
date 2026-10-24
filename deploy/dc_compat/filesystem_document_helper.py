# Copyright 2026 TrillionniumFoundation; SPDX-License-Identifier: Apache-2.0
"""Isolated independent Office/PDF parsers; JSON request/response over stdio.

DOCX and existing OOXML spreadsheet edits retain every unrelated ZIP member.
No Desktop Commander or vendor Node module is imported or executed.
"""
from __future__ import annotations
import base64
import copy
import datetime
import io
import json
import os
import pathlib
import re
import shutil
import stat
import sys
import tempfile
import urllib.parse
import zipfile
from xml.sax.saxutils import escape


def atomic(path, data):
    path = pathlib.Path(path)
    fd, temporary = tempfile.mkstemp(prefix='.hepta-document-', dir=path.parent)
    try:
        if path.exists(): os.fchmod(fd, stat.S_IMODE(path.stat().st_mode))
        else: os.fchmod(fd, 0o644)
        with os.fdopen(fd, 'wb') as output:
            output.write(data); output.flush(); os.fsync(output.fileno())
        if path.exists():
            shutil.copystat(path, temporary)
            os.utime(temporary, None)
        os.replace(temporary, path)
    finally:
        try: os.unlink(temporary)
        except FileNotFoundError: pass


def pretty(xml):
    lines, depth = [], 0
    for part in re.split(r'(?<=>)(?=<)', xml):
        part = part.strip()
        if not part: continue
        closing, single, instruction = part.startswith('</'), part.endswith('/>'), part.startswith('<?')
        inline = not closing and not single and '</' in part
        if closing: depth = max(0, depth - 1)
        lines.append('  ' * depth + part)
        if not closing and not single and not inline and not instruction: depth += 1
    return '\n'.join(lines)


def xml_root(raw):
    from lxml import etree
    return etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))


def zip_rewrite(path, changes):
    source, target = io.BytesIO(pathlib.Path(path).read_bytes()), io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, 'w') as rewritten:
        for info in original.infolist():
            rewritten.writestr(info, changes.pop(info.filename, original.read(info.filename)))
        for name, data in changes.items(): rewritten.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)
    atomic(path, target.getvalue())


def docx_read(path, args):
    with zipfile.ZipFile(path) as archive:
        raw = archive.read('word/document.xml').decode('utf-8')
        formatted = pretty(raw); lines = formatted.split('\n')
        offset, length = int(args.get('offset') or 0), args.get('length')
        if offset:
            start = max(0, len(lines) + offset) if offset < 0 else offset
            selected = lines[start:] if offset < 0 or length is None else lines[start:start + int(length)]
            return {'kind': 'text', 'text': f'[DOCX XML: lines {start}-{start + len(selected) - 1} of {len(lines)} ({max(0, len(lines)-start-len(selected))} remaining)]\n' + '\n'.join(selected)}
        root = xml_root(raw.encode()); ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        body = root.find('w:body', ns)
        outline = []
        if body is not None:
            for index, element in enumerate(body):
                local = element.tag.rsplit('}', 1)[-1]
                if local == 'tbl':
                    rows = element.findall('w:tr', ns)
                    text = f'[{index}] w:tbl ({len(rows)} rows)'
                    for number, row in enumerate(rows):
                        cells = [''.join(cell.itertext()) for cell in row.findall('w:tc', ns)]
                        text += f'\n  row{number}: [' + ' | '.join(cells) + ']'
                else:
                    fragments = re.findall(r'<w:t(?:\s[^>]*)?>[^<]*</w:t>', raw_fragment(element))
                    text = f'[{index}] w:{local}'
                    style = element.find('.//w:pStyle', ns)
                    if style is not None: text += ' style="' + style.get('{' + ns['w'] + '}val', '') + '"'
                    if fragments: text += '\n  ' + ''.join(fragments)
                    elif local == 'p': text += ' (empty)' if not element.findall('.//w:drawing', ns) else ' [drawing/image]'
                outline.append(text)
        for name in archive.namelist():
            if re.fullmatch(r'word/(?:header|footer)\d+\.xml', name):
                content = archive.read(name).decode()
                outline.append(name + '\n  ' + ''.join(re.findall(r'<w:t(?:\s[^>]*)?>[^<]*</w:t>', content)))
        return {'kind': 'text', 'text': f'DOCX Outline: {len(body) if body is not None else 0} body children\nEdit XML with edit_block(old_string="<w:t>old text</w:t>", new_string="<w:t>new text</w:t>").\nRaw XML: use nonzero offset for pretty-printed XML.\n' + '\n'.join(outline) + f'\n\nRaw XML: {len(lines)} lines.'}


def raw_fragment(element):
    from lxml import etree
    return etree.tostring(element, encoding='unicode')


def docx_write(path, args):
    if args.get('mode') == 'append': raise ValueError('DOCX append not supported. Use edit_block to modify an existing DOCX.')
    from docx import Document
    document = Document()
    for line in args['content'].split('\n'):
        heading = re.match(r'^(#{1,6})\s+(.+)', line)
        if heading: document.add_heading(heading[2], level=len(heading[1]))
        else: document.add_paragraph(line)
    output = io.BytesIO(); document.save(output); atomic(path, output.getvalue())


def docx_edit(path, args):
    old, new = args.get('old_string'), args.get('new_string')
    if not isinstance(old, str) or not old or not isinstance(new, str): raise ValueError('DOCX editing requires old_string and new_string')
    expected = int(args.get('expected_replacements', 1))
    with zipfile.ZipFile(path) as archive:
        candidates = ['word/document.xml'] + [f'word/{kind}{number}.xml' for kind in ('header', 'footer') for number in (1, 2, 3)]
        for name in candidates:
            if name not in archive.namelist(): continue
            formatted = pretty(archive.read(name).decode())
            count = formatted.count(old)
            if not count: continue
            if count != expected: raise ValueError(f'Expected {expected} occurrence(s) but found {count}; DOCX was not modified')
            changed = ''.join(line.lstrip() for line in formatted.replace(old, new).split('\n'))
            xml_root(changed.encode())  # Refuse malformed output before publication.
            zip_rewrite(path, {name: changed.encode()})
            return True
    raise ValueError('Search string not found in DOCX; file was not modified')


def excel_load(path, data_only=False):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(pathlib.Path(path).read_bytes()), data_only=data_only,
                                 keep_vba=str(path).lower().endswith('.xlsm'), rich_text=True)


def excel_sheets(path):
    if not zipfile.is_zipfile(path):
        import xlrd
        book = xlrd.open_workbook(path)
        return [{'name': sheet.name, 'rowCount': sheet.nrows, 'colCount': sheet.ncols} for sheet in book.sheets()]
    book = excel_load(path)
    result = [{'name': sheet.title, 'rowCount': sheet.max_row if any(cell.value is not None for row in sheet for cell in row) else 0,
               'colCount': sheet.max_column if any(cell.value is not None for row in sheet for cell in row) else 0} for sheet in book]
    book.close(); return result


def excel_read(path, args):
    from openpyxl.utils.cell import range_boundaries
    sheet_ref, region = args.get('sheet'), args.get('range')
    if region and '!' in region:
        embedded, region = region.rsplit('!', 1)
        if sheet_ref is None: sheet_ref = embedded.strip("'").replace("''", "'")
    if isinstance(sheet_ref, str) and sheet_ref.isdigit(): sheet_ref = int(sheet_ref)
    legacy = not zipfile.is_zipfile(path)
    if legacy:
        import xlrd
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(sheet_ref or 0) if isinstance(sheet_ref, (int, type(None))) else book.sheet_by_name(sheet_ref)
        name, total, columns = sheet.name, sheet.nrows, sheet.ncols
        value_at = lambda row, col: sheet.cell_value(row - 1, col - 1) if row <= sheet.nrows and col <= sheet.ncols else None
    else:
        book = excel_load(path, data_only=True)
        sheet = book.worksheets[sheet_ref or 0] if isinstance(sheet_ref, (int, type(None))) else book[sheet_ref]
        name, total, columns = sheet.title, sheet.max_row, sheet.max_column
        # openpyxl collapses an empty cached formula string into None. Preserve
        # that distinction, and ExcelJS's structured cached error projection.
        overrides = {}
        with zipfile.ZipFile(path) as archive:
            ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            workbook = xml_root(archive.read('xl/workbook.xml'))
            relationships = xml_root(archive.read('xl/_rels/workbook.xml.rels'))
            entry = next(item for item in workbook.find('{' + ns + '}sheets') if item.get('name') == name)
            identity = entry.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = next(item.get('Target') for item in relationships if item.get('Id') == identity)
            member = target.lstrip('/') if target.startswith('/') else str(pathlib.PurePosixPath('xl') / target)
            root = xml_root(archive.read(member))
            for cell in root.iter('{' + ns + '}c'):
                if cell.find('{' + ns + '}f') is None: continue
                cached = cell.find('{' + ns + '}v')
                if cached is not None and cell.get('t') == 'str' and cached.text is None: overrides[cell.get('r')] = ''
                if cached is not None and cell.get('t') == 'e': overrides[cell.get('r')] = {'error': cached.text or ''}
        value_at = lambda row, col: overrides.get(sheet.cell(row, col).coordinate, sheet.cell(row, col).value)
    start_col, start_row, end_col, end_row = range_boundaries(region) if region else (1, 1, max(1, columns), max(1, total))
    total_rows = end_row - start_row + 1
    offset, length = int(args.get('offset') or 0), args.get('length', 1000)
    if offset < 0: start_row = max(start_row, end_row + offset + 1)
    else:
        start_row += offset
        if length is not None and int(length) > 0: end_row = min(end_row, start_row + int(length) - 1)
    def value(row, col):
        cell = value_at(row, col)
        if isinstance(cell, (datetime.date, datetime.datetime, datetime.time)): return cell.isoformat()
        if isinstance(cell, dict): return cell
        if isinstance(cell, (str, int, float, bool, type(None))): return cell
        return str(cell)
    data = [[value(row, col) for col in range(start_col, end_col + 1)] for row in range(start_row, end_row + 1)]
    if not legacy: book.close()
    text = f"[Sheet: '{name}' from {path}]"
    if len(data) < total_rows: text += f'\n[Showing rows {start_row}-{start_row + len(data) - 1} of {total_rows} total. Use offset/length to paginate.]'
    return {'kind': 'text', 'text': text + '\n[To MODIFY cells: use edit_block with range, e.g. Sheet1!A1:B2.]\n\n' + json.dumps(data, ensure_ascii=False)}


def _excel_rich_font(value, namespace):
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles.colors import Color
    names = {'name': 'rFont', 'charset': 'charset', 'family': 'family', 'bold': 'b', 'italic': 'i',
             'strike': 'strike', 'outline': 'outline', 'shadow': 'shadow', 'condense': 'condense',
             'extend': 'extend', 'size': 'sz', 'underline': 'u', 'vertAlign': 'vertAlign', 'scheme': 'scheme'}
    if not isinstance(value, dict) or set(value) - set(names) - {'color'}:
        raise ValueError('Unsupported richText font properties; workbook was not modified')
    mapped = {names[key]: val for key, val in value.items() if key in names}
    if 'underline' in value and isinstance(value['underline'], bool): mapped['u'] = 'single' if value['underline'] else None
    if 'color' in value:
        color = value['color']
        if not isinstance(color, dict) or set(color) - {'argb', 'theme', 'tint', 'indexed', 'auto'}:
            raise ValueError('Unsupported richText font color; workbook was not modified')
        if color.get('argb'): mapped['color'] = Color(rgb=color['argb'])
        elif color.get('theme') is not None: mapped['color'] = Color(theme=color['theme'], tint=color.get('tint', 0))
        elif color.get('indexed') is not None: mapped['color'] = Color(indexed=color['indexed'])
        else: mapped['color'] = Color(auto=True)
    root = InlineFont(**mapped).to_tree(tagname='rPr')
    for element in root.iter(): element.tag = '{' + namespace + '}' + element.tag.rsplit('}', 1)[-1]
    return root


def _excel_formula(cell, value, namespace, worksheet):
    from lxml import etree
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    formula = etree.SubElement(cell, '{' + namespace + '}f')
    if value.get('formula'):
        if not isinstance(value['formula'], str): raise ValueError('Excel formula must be a string')
        formula.text = value['formula'].lstrip('=')
        kind = value.get('shareType')
        if kind not in (None, 'shared', 'array'): raise ValueError('Unsupported Excel formula shareType')
        if kind:
            formula.set('t', kind)
            reference = value.get('ref', cell.get('r'))
            range_boundaries(reference)
            formula.set('ref', reference)
            if kind == 'shared':
                used = [int(item.get('si')) for item in worksheet.iter('{' + namespace + '}f') if item.get('si') is not None]
                formula.set('si', str(max(used, default=-1) + 1))
    else:
        reference = value.get('sharedFormula')
        if not isinstance(reference, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', reference):
            raise ValueError('sharedFormula must identify a master cell')
        master = worksheet.find('.//{' + namespace + '}c[@r="' + reference.replace('"', '') + '"]/{' + namespace + '}f')
        if master is None or not master.text:
            raise ValueError('Shared formula master must already exist; workbook was not modified')
        address = cell.get('r')
        if master.get('t') not in (None, 'shared'): raise ValueError('Shared formula master is an array formula')
        if master.get('si') is None:
            used = [int(item.get('si')) for item in worksheet.iter('{' + namespace + '}f') if item.get('si') is not None]
            master.set('si', str(max(used, default=-1) + 1))
        master.set('t', 'shared')
        left, top, right, bottom = range_boundaries(master.get('ref', reference))
        col, row, _, _ = range_boundaries(address)
        master.set('ref', f'{get_column_letter(min(left, col))}{min(top, row)}:{get_column_letter(max(right, col))}{max(bottom, row)}')
        formula.set('t', 'shared'); formula.set('si', master.get('si'))
    result = value.get('result')
    if result is None: return
    cached = etree.SubElement(cell, '{' + namespace + '}v')
    if isinstance(result, bool): cell.set('t', 'b'); cached.text = str(int(result))
    elif isinstance(result, (int, float)): cached.text = json.dumps(result, allow_nan=False)
    elif isinstance(result, str): cell.set('t', 'str'); cached.text = result
    elif isinstance(result, dict) and isinstance(result.get('error'), str): cell.set('t', 'e'); cached.text = result['error']
    else: raise ValueError('Unsupported cached formula result; workbook was not modified')


def _excel_cell(cell, value, namespace, worksheet=None):
    from lxml import etree
    for child in list(cell):
        if child.tag.rsplit('}', 1)[-1] in ('v', 'f', 'is'): cell.remove(child)
    cell.attrib.pop('t', None)
    def element(name): return etree.SubElement(cell, '{' + namespace + '}' + name)
    if value is None: return
    if isinstance(value, dict) and value.get('text') and value.get('hyperlink'):
        if not isinstance(value['text'], str) or not isinstance(value['hyperlink'], str):
            raise ValueError('Hyperlink text and target must be strings')
        cell.set('t', 'inlineStr'); text = etree.SubElement(element('is'), '{' + namespace + '}t')
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); text.text = value['text']
        return value
    if isinstance(value, dict) and (value.get('formula') or value.get('sharedFormula')):
        _excel_formula(cell, value, namespace, worksheet)
    elif isinstance(value, str) and value.startswith('='):
        element('f').text = value[1:]
        element('v').text = ''  # Keep the established explicit empty-cache representation for bare formula strings.
    elif isinstance(value, bool): cell.set('t', 'b'); element('v').text = str(int(value))
    elif isinstance(value, (int, float)): element('v').text = json.dumps(value, allow_nan=False)
    elif isinstance(value, dict) and 'richText' in value and (isinstance(value['richText'], (dict, list)) or bool(value['richText'])):
        runs = value['richText']
        if not isinstance(runs, list): raise ValueError('richText must be an array of formatted text runs')
        cell.set('t', 'inlineStr'); inline = element('is')
        for run in runs:
            if not isinstance(run, dict) or not isinstance(run.get('text'), str) or set(run) - {'text', 'font'}:
                raise ValueError('richText runs must contain text and optional font; workbook was not modified')
            node = etree.SubElement(inline, '{' + namespace + '}r')
            if run.get('font') is not None: node.append(_excel_rich_font(run['font'], namespace))
            text = etree.SubElement(node, '{' + namespace + '}t')
            text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); text.text = run['text']
    elif isinstance(value, dict) and value.get('error'):
        if not isinstance(value['error'], str): raise ValueError('Excel error must be a string')
        cell.set('t', 'e'); element('v').text = value['error']
    else:
        if isinstance(value, dict) and value.get('sharedString'):
            raise ValueError('External sharedString indexes are not accepted without their source table')
        if isinstance(value, (dict, list)): value = json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
        cell.set('t', 'inlineStr'); text = etree.SubElement(element('is'), '{' + namespace + '}t')
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); text.text = str(value)


class _ExcelHyperlinks:
    def __init__(self, worksheet, relationships, namespace):
        self.worksheet, self.relationships, self.namespace = worksheet, relationships, namespace
        self.dirty = False

    def set(self, address, value):
        from lxml import etree
        q = '{' + self.namespace + '}'
        links = self.worksheet.find(q + 'hyperlinks')
        relation_namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        id_key = '{' + relation_namespace + '}id'
        if links is not None:
            for link in list(links):
                if link.get('ref') != address: continue
                previous = link.get(id_key); links.remove(link); self.dirty = True
                if previous and not any(element.get(id_key) == previous for element in self.worksheet.iter()):
                    for relationship in list(self.relationships):
                        if relationship.get('Id') == previous and relationship.get('Type', '').endswith('/hyperlink'):
                            self.relationships.remove(relationship)
        if value is None:
            if links is not None and not len(links): self.worksheet.remove(links)
            return
        if links is None:
            links = etree.Element(q + 'hyperlinks')
            following = {'printOptions', 'pageMargins', 'pageSetup', 'headerFooter', 'rowBreaks', 'colBreaks',
                         'customProperties', 'cellWatches', 'ignoredErrors', 'smartTags', 'drawing', 'legacyDrawing',
                         'legacyDrawingHF', 'picture', 'oleObjects', 'controls', 'webPublishItems', 'tableParts', 'extLst'}
            index = next((number for number, child in enumerate(self.worksheet) if child.tag.rsplit('}', 1)[-1] in following), len(self.worksheet))
            self.worksheet.insert(index, links)
        used = {relationship.get('Id') for relationship in self.relationships}
        index = 1
        while 'rId' + str(index) in used: index += 1
        identity = 'rId' + str(index)
        etree.SubElement(self.relationships, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                         Id=identity, Type=relation_namespace + '/hyperlink', Target=value['hyperlink'], TargetMode='External')
        link = etree.SubElement(links, q + 'hyperlink', ref=address)
        link.set(id_key, identity)
        if value.get('tooltip') is not None: link.set('tooltip', str(value['tooltip']))
        if re.fullmatch(r'[^!]+![a-zA-Z]+[0-9]+', value['hyperlink']): link.set('location', value['hyperlink'])
        self.dirty = True


def excel_patch(path, sheet_name, matrix, cell_range=None, append=False):
    from lxml import etree
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    if not zipfile.is_zipfile(path): raise ValueError('In-place legacy BIFF editing cannot preserve its original workbook; convert to OOXML explicitly first')
    if not isinstance(matrix, list) or any(not isinstance(row, list) for row in matrix): raise ValueError('Excel content must be a 2D array')
    with zipfile.ZipFile(path) as archive:
        workbook = xml_root(archive.read('xl/workbook.xml')); rels = xml_root(archive.read('xl/_rels/workbook.xml.rels'))
        namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        sheets = workbook.find('{' + namespace + '}sheets')
        entry = next((sheet for sheet in sheets if sheet.get('name') == sheet_name), None)
        changes = {}
        if entry is None:
            if not sheet_name or len(sheet_name) > 31 or re.search(r'[\\/*?:\[\]]', sheet_name): raise ValueError('Invalid Excel sheet name')
            sheet_id = max((int(sheet.get('sheetId')) for sheet in sheets), default=0) + 1
            used_ids = {relationship.get('Id') for relationship in rels}
            number = 1
            while 'rId' + str(number) in used_ids: number += 1
            relationship_id = 'rId' + str(number)
            part_number = 1
            while f'xl/worksheets/sheet{part_number}.xml' in archive.namelist(): part_number += 1
            target = f'worksheets/sheet{part_number}.xml'
            entry = etree.SubElement(sheets, '{' + namespace + '}sheet', name=sheet_name, sheetId=str(sheet_id))
            entry.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id', relationship_id)
            etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                Id=relationship_id, Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet', Target=target)
            types = xml_root(archive.read('[Content_Types].xml'))
            etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                PartName='/xl/' + target, ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
            changes.update({'xl/workbook.xml': etree.tostring(workbook),
                'xl/_rels/workbook.xml.rels': etree.tostring(rels), '[Content_Types].xml': etree.tostring(types)})
        relationship = entry.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        target = next(rel.get('Target') for rel in rels if rel.get('Id') == relationship)
        member = target.lstrip('/') if target.startswith('/') else str(pathlib.PurePosixPath('xl') / target)
        if member in archive.namelist(): root = xml_root(archive.read(member))
        else:
            root = etree.Element('{' + namespace + '}worksheet', nsmap={None: namespace})
            etree.SubElement(root, '{' + namespace + '}dimension', ref='A1')
            etree.SubElement(root, '{' + namespace + '}sheetData')
        relative_member = pathlib.PurePosixPath(member)
        relationship_member = str(relative_member.parent / '_rels' / (relative_member.name + '.rels'))
        worksheet_relationships = (xml_root(archive.read(relationship_member)) if relationship_member in archive.namelist()
                                  else etree.Element('{http://schemas.openxmlformats.org/package/2006/relationships}Relationships',
                                                     nsmap={None: 'http://schemas.openxmlformats.org/package/2006/relationships'}))
        hyperlinks = _ExcelHyperlinks(root, worksheet_relationships, namespace)
        data = root.find('{' + namespace + '}sheetData')
        used = [(int(row.get('r')), row) for row in data]
        start_col, start_row = 1, 1
        if append: start_row = max((number for number, row in used if any(child.tag.endswith('}v') or child.tag.endswith('}is') or child.tag.endswith('}f') for cell in row for child in cell)), default=0) + 1
        elif cell_range: start_col, start_row, _, _ = range_boundaries(cell_range)
        else:
            for row in data:
                for cell in row:
                    _excel_cell(cell, None, namespace, root)
                    hyperlinks.set(cell.get('r'), None)
        rows = {int(row.get('r')): row for row in data}
        for r, row_data in enumerate(matrix, start_row):
            row = rows.get(r)
            if row is None:
                row = etree.Element('{' + namespace + '}row', r=str(r)); data.append(row); rows[r] = row
            cells = {cell.get('r'): cell for cell in row}
            for c, value in enumerate(row_data, start_col):
                address = get_column_letter(c) + str(r); cell = cells.get(address)
                if cell is None: cell = etree.SubElement(row, '{' + namespace + '}c', r=address)
                link = _excel_cell(cell, value, namespace, root)
                hyperlinks.set(address, link)
            row[:] = sorted(row, key=lambda cell: range_boundaries(cell.get('r'))[0])
        data[:] = sorted(data, key=lambda row: int(row.get('r')))
        shared_masters = {formula.get('si') for formula in root.iter('{' + namespace + '}f')
                          if formula.get('t') == 'shared' and formula.text}
        if any(formula.get('t') == 'shared' and not formula.text and formula.get('si') not in shared_masters
               for formula in root.iter('{' + namespace + '}f')):
            raise ValueError('Edit would orphan an existing shared formula; workbook was not modified')
        dimension = root.find('{' + namespace + '}dimension')
        if dimension is not None:
            last_row = max(rows, default=1); last_col = max((range_boundaries(cell.get('r'))[0] for row in data for cell in row), default=1)
            dimension.set('ref', f'A1:{get_column_letter(last_col)}{last_row}')
        changed = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        if hyperlinks.dirty:
            changes[relationship_member] = etree.tostring(worksheet_relationships, xml_declaration=True, encoding='UTF-8', standalone=True)
    changes[member] = changed
    zip_rewrite(path, changes)


def excel_write(path, args):
    import openpyxl
    content = json.loads(args['content']) if isinstance(args['content'], str) else args['content']
    sheets = {'Sheet1': content} if isinstance(content, list) else content
    if not isinstance(sheets, dict) or any(not isinstance(rows, list) for rows in sheets.values()): raise ValueError('Expected a 2D array or object mapping sheet names to 2D arrays')
    if args.get('mode') == 'append' and pathlib.Path(path).exists():
        with tempfile.TemporaryDirectory(prefix='.hepta-excel-append-', dir=pathlib.Path(path).parent) as directory:
            staging = pathlib.Path(directory) / pathlib.Path(path).name
            shutil.copy2(path, staging)
            for name, matrix in sheets.items(): excel_patch(staging, name, matrix, append=True)
            atomic(path, staging.read_bytes())
        return True
    # Build structural/style parts with openpyxl, then write all cell values
    # through the same OOXML encoder used for edits. An openpyxl save after
    # setting formulas would otherwise discard their supplied cached results.
    with tempfile.TemporaryDirectory(prefix='.hepta-excel-write-', dir=pathlib.Path(path).parent) as directory:
        staging = pathlib.Path(directory) / pathlib.Path(path).name
        workbook = openpyxl.Workbook(); workbook.remove(workbook.active)
        for name, matrix in sheets.items():
            if any(not isinstance(row, list) for row in matrix): raise ValueError('Expected a 2D array')
            workbook.create_sheet(name)
        if not workbook.worksheets: workbook.create_sheet('Sheet1')
        workbook.save(staging); workbook.close()
        for name, matrix in sheets.items(): excel_patch(staging, name, matrix)
        atomic(path, staging.read_bytes())
    return True


def _pdf_markdown(markdown, options, request):
    # -I intentionally excludes the helper's directory from sys.path. Load
    # this exact trusted sibling without shadowing standard-library modules.
    import importlib.util
    spec = importlib.util.spec_from_file_location('_hepta_private_chromium_pdf', pathlib.Path(__file__).with_name('chromium_pdf.py'))
    renderer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(renderer)
    return renderer.render_markdown(markdown, options, request)


def _pdf_original_layout(reader):
    if not reader.pages: return {}
    page = reader.pages[0]
    media, crop = page.mediabox, page.cropbox
    margins = {'left': float(crop.left - media.left), 'bottom': float(crop.bottom - media.bottom),
               'right': float(media.right - crop.right), 'top': float(media.top - crop.top)}
    if not any(margins.values()): margins = dict.fromkeys(margins, 72)
    return {'format': None, 'width': f'{float(media.width) / 72:.4f}in', 'height': f'{float(media.height) / 72:.4f}in',
            'margin': {side: f'{points / 72:.4f}in' for side, points in margins.items()}}


def pdf_read(path, args):
    from pypdf import PdfReader
    reader = PdfReader(path); count = len(reader.pages); metadata = reader.metadata or {}
    offset, length = int(args.get('offset') or 0), args.get('length')
    start = max(0, count + offset) if offset < 0 else offset
    finish = count if offset < 0 or length is None or int(length) == 0 else min(count, start + int(length))
    pages = []
    for number in range(start, finish):
        page = reader.pages[number]; images = []
        for image in page.images:
            mime = 'image/png' if image.data.startswith(b'\x89PNG') else 'image/jpeg' if image.data.startswith(b'\xff\xd8') else None
            if mime is None:
                from PIL import Image
                buffer = io.BytesIO(); Image.open(io.BytesIO(image.data)).save(buffer, 'PNG'); raw = buffer.getvalue(); mime = 'image/png'
            else: raw = image.data
            images.append({'type': 'image', 'data': base64.b64encode(raw).decode(), 'mimeType': mime})
        pages.append({'pageNumber': number + 1, 'text': page.extract_text() or '', 'images': images})
    return {'kind': 'pdf', 'totalPages': count, 'title': str(metadata.get('/Title', '')), 'author': str(metadata.get('/Author', '')), 'pages': pages}


def pdf_write(path, args, request):
    from pypdf import PdfReader, PdfWriter
    content = args['content']; target = args.get('outputPath') or path
    if isinstance(content, str): atomic(target, _pdf_markdown(content, args.get('options'), request)); return True
    if not isinstance(content, list): raise ValueError('PDF content must be Markdown or an operation array')
    original = PdfReader(path)
    layout = _pdf_original_layout(original)
    writer = PdfWriter(); writer.clone_document_from_reader(original)
    for operation in content:
        if operation['type'] == 'delete':
            count = len(writer.pages)
            indices = {index if index >= 0 else count + index for index in operation['pageIndexes']}
            for index in sorted((index for index in indices if 0 <= index < count), reverse=True): writer.remove_page(index)
        elif operation['type'] == 'insert':
            if 'markdown' in operation:
                supplied = operation.get('pdfOptions') or {}
                if not isinstance(supplied, dict): raise ValueError('pdfOptions must be an object')
                supplied = dict(supplied)
                supplied['pdf_options'] = dict(layout, **supplied.get('pdf_options', {}))
                # Flat PDFOptions remain supported; explicit supplied options
                # win over the inherited original-page layout.
                for key in layout:
                    if key in supplied: supplied['pdf_options'][key] = supplied.pop(key)
                source = PdfReader(io.BytesIO(_pdf_markdown(operation['markdown'], supplied, request)))
            elif operation.get('sourcePdfPath'): source = PdfReader(operation['sourcePdfPath'])
            else: raise ValueError('No source provided for PDF insert operation')
            position = int(operation['pageIndex']); position = len(writer.pages) + position if position < 0 else position
            if not 0 <= position <= len(writer.pages): raise ValueError('Invalid page index')
            for index, page in enumerate(source.pages): writer.insert_page(page, position + index)
        else: raise ValueError('Unknown PDF operation')
    output = io.BytesIO(); writer.write(output); atomic(target, output.getvalue()); return True


def invoke(request):
    operation, path, args = request['operation'], pathlib.Path(request['path']), request.get('arguments', {})
    suffix = path.suffix.lower()
    if operation == 'pdf': return pdf_write(path, args, request)
    if operation == 'read':
        if suffix == '.docx': return docx_read(path, args)
        if suffix == '.pdf': return pdf_read(path, args)
        return excel_read(path, args)
    if operation == 'write':
        if suffix == '.docx': return docx_write(path, args)
        if suffix == '.pdf': return pdf_write(path, args, request)
        return excel_write(path, args)
    if operation == 'edit':
        if suffix == '.docx': return docx_edit(path, args)
        if suffix == '.pdf': return pdf_write(path, args, request)
        region = args.get('range')
        if not isinstance(region, str) or not region: raise ValueError('Excel editing requires a range')
        sheet, cells = region.rsplit('!', 1) if '!' in region else (region, None)
        return excel_patch(path, sheet.strip("'").replace("''", "'"), args['content'], cells)
    if operation == 'info':
        if suffix in ('.xlsx', '.xlsm', '.xls'): return {'fileType': 'excel', 'isExcelFile': True, 'sheets': excel_sheets(path)}
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path); metadata = reader.metadata or {}
            return {'fileType': 'binary', 'isPdf': True, 'totalPages': len(reader.pages), 'title': str(metadata.get('/Title', '')), 'author': str(metadata.get('/Author', ''))}
        with zipfile.ZipFile(path) as archive: count = len(pretty(archive.read('word/document.xml').decode()).split('\n'))
        return {'fileType': 'docx', 'lineCount': count, 'lastLine': count - 1, 'appendPosition': count}
    raise ValueError('Unknown document operation')


if __name__ == '__main__':
    try:
        result = invoke(json.load(sys.stdin))
        response = {'ok': True, 'value': result}
    except Exception as error:
        response = {'ok': False, 'error': str(error), 'error_type': type(error).__name__}
    sys.stdout.write(json.dumps(response, ensure_ascii=False, default=str))
